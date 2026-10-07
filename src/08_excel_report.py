# Step 8: one Excel report with every KPI table and chart, built automatically
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

OUTPUT = Path("output")
REPORT = OUTPUT / "olist_delivery_report.xlsx"
NAVY = "1F3864"

# 1. Collect the tables made in steps 2-6
overall = pd.read_csv(OUTPUT / "kpi_overall.csv")
network = pd.read_csv(OUTPUT / "kpi_dpmo_network.csv")
seller_kpis = pd.DataFrame([
    {"kpi": "Network DPMO (late + bad review)", "value": network["network_dpmo"].iloc[0]},
    {"kpi": "Audit threshold (DPMO)", "value": network["audit_threshold"].iloc[0]},
    {"kpi": "Sellers with 30+ orders", "value": network["sellers_active"].iloc[0]},
    {"kpi": "Sellers flagged for audit", "value": network["sellers_flagged"].iloc[0]},
])
summary = pd.concat([overall, seller_kpis], ignore_index=True)

audit = pd.read_csv(OUTPUT / "audit_list.csv")
audit.insert(0, "priority", range(1, len(audit) + 1))
audit = audit[["priority", "seller_id", "orders", "late", "bad_reviews", "dpmo", "excess_defects"]]

sheets = {
    "Summary": summary,
    "Monthly": pd.read_csv(OUTPUT / "kpi_monthly.csv"),
    "States": pd.read_csv(OUTPUT / "kpi_by_state.csv"),
    "Reviews": pd.read_csv(OUTPUT / "kpi_reviews_by_delay.csv"),
    "Audit list": audit,
    "Data quality": pd.read_csv(OUTPUT / "quality_log.csv"),
}
titles = {
    "Summary": "Delivery KPIs - summary",
    "Monthly": "Late deliveries by month of purchase",
    "States": "Late deliveries by customer state (min. 500 orders)",
    "Reviews": "Review score by delivery delay",
    "Audit list": "Sellers flagged for audit (sorted by excess defects)",
    "Data quality": "Data quality checks",
}
charts = {
    "Monthly": "01_late_rate_by_month.png",
    "States": "02_late_rate_by_state.png",
    "Reviews": "03_review_by_delay.png",
    "Audit list": "04_seller_dpmo.png",
}

# 2. Write every table to its own sheet, starting on row 3 (rows 1-2 for the title)
with pd.ExcelWriter(REPORT, engine="openpyxl") as writer:
    for sheet_name, table in sheets.items():
        table.to_excel(writer, sheet_name=sheet_name, index=False, startrow=2)

# 3. Open the file again and format it
wb = load_workbook(REPORT)
for ws in wb.worksheets:
    ws["A1"] = titles[ws.title]
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    ws["A2"] = f"Created automatically on {date.today()} | Source: Olist public dataset (Kaggle)"
    ws["A2"].font = Font(italic=True, color="808080")

    for cell in ws[3]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", start_color=NAVY)

    for column_cells in ws.iter_cols(min_row=3):
        longest = max((len(str(c.value)) for c in column_cells if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = longest + 3

    ws.freeze_panes = "A4"

    # Thousands separator for big numbers (96447 -> 96,447)
    for row in ws.iter_rows(min_row=4):
        for cell in row:
            if isinstance(cell.value, (int, float)) and abs(cell.value) >= 1000:
                cell.number_format = "#,##0"

# 4. Put each chart next to its table
for sheet_name, file_name in charts.items():
    ws = wb[sheet_name]
    picture = Image(OUTPUT / "charts" / file_name)
    picture.width = picture.width * 0.5
    picture.height = picture.height * 0.5
    ws.add_image(picture, f"{get_column_letter(ws.max_column + 2)}3")

wb.save(REPORT)
print(f"Saved: {REPORT}")
